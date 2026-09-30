"""Independent data-only XLSX fixture/validation, executed only in cloud CI."""
import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('mode', choices=['generate', 'verify'])
    parser.add_argument('directory', type=Path)
    args = parser.parse_args()
    if args.mode == 'generate':
        args.directory.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '売上'
        sheet['A1'] = '商品'
        sheet['B2'] = 12.5
        sheet['B3'] = '=B2*2'
        sheet['C1'] = '00123'
        sheet['D1'] = '=literal'
        sheet['D1'].data_type = 's'
        workbook.create_sheet('予算')['A1'] = ' 日本語 <&> '
        workbook.save(args.directory / 'input.xlsx')
    else:
        workbook = load_workbook(args.directory / 'output.xlsx', data_only=False)
        assert workbook.sheetnames == ['売上', '予算']
        sheet = workbook['売上']
        assert sheet['A1'].value == '商品'
        assert sheet['B2'].value == 12.5 and sheet['B2'].data_type == 'n'
        assert sheet['B3'].value == '=B2*2' and sheet['B3'].data_type == 'f'
        assert sheet['C1'].value == '00123' and sheet['C1'].data_type == 's'
        assert sheet['D1'].value == '=literal' and sheet['D1'].data_type == 's'
        assert workbook['予算']['A1'].value == ' 日本語 <&> '
        print('Independent openpyxl XLSX compatibility: PASS')


if __name__ == '__main__':
    main()
